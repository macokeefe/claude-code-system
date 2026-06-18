#!/usr/bin/env python3
# Build an editable PowerPoint of the Meritage line-balance findings.
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

NAVY=RGBColor.from_string('15171A'); INK=RGBColor.from_string('1F2326')
GOLD=RGBColor.from_string('A8854C'); GOLDD=RGBColor.from_string('8A6B39')
RUST=RGBColor.from_string('B1542F'); GOOD=RGBColor.from_string('4F7A5B')
MUTED=RGBColor.from_string('6B7178'); LINE=RGBColor.from_string('E0DDD6')
WHITE=RGBColor.from_string('FFFFFF'); PAPER=RGBColor.from_string('FBFAF8')
BODY='Calibri'; SERIF='Times New Roman'

COL={2:'A8854C',3:'4F7A5B',4:'15171A',5:'3C4A52',6:'8A6B39',7:'6B7178',8:'B1542F',9:'2F6B5A',10:'7A5230',11:'9A7B4F',12:'C9A55F'}
STEPS=[(2,'Connector Prep',19.25),(3,'Leg Assembly',8),(4,'Arms Assembly',64),(5,'Seat Frame',47),
       (6,'Back Frame',43.5),(7,'Attaching Arms',18),(8,'Back Support',5),(9,'Leg Finishing',8),
       (10,'Seat Support',22.5),(11,'Attaching Seat Frame',18),(12,'TUUCI Plate',1)]
TAKT=42.0

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
SW,SH=13.333,7.5
BLANK=prs.slide_layouts[6]

def slide():
    s=prs.slides.add_slide(BLANK)
    bg=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,prs.slide_width,prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb=PAPER; bg.line.fill.background(); bg.shadow.inherit=False
    return s

def noline(sp): sp.line.fill.background()
def noshadow(sp): sp.shadow.inherit=False

def text(s,x,y,w,h,runs,size=14,color=INK,bold=False,align='l',font=BODY,anchor=None,sp_after=4,line_spacing=None):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=tb.text_frame; tf.word_wrap=True
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    if anchor is not None: tf.vertical_anchor=anchor
    if isinstance(runs,str): runs=[(runs,size,color,bold,font)]
    p=tf.paragraphs[0]; p.alignment={'l':PP_ALIGN.LEFT,'c':PP_ALIGN.CENTER,'r':PP_ALIGN.RIGHT}[align]
    if line_spacing: p.line_spacing=line_spacing
    for t,sz,c,b,fn in runs:
        r=p.add_run(); r.text=t; r.font.size=Pt(sz); r.font.bold=b; r.font.name=fn; r.font.color.rgb=c
    return tb

def para(tb,runs,align='l',space_before=6,line_spacing=1.1):
    tf=tb.text_frame; p=tf.add_paragraph(); p.alignment={'l':PP_ALIGN.LEFT,'c':PP_ALIGN.CENTER,'r':PP_ALIGN.RIGHT}[align]
    p.space_before=Pt(space_before); p.line_spacing=line_spacing
    for t,sz,c,b,fn in runs:
        r=p.add_run(); r.text=t; r.font.size=Pt(sz); r.font.bold=b; r.font.name=fn; r.font.color.rgb=c
    return p

def rect(s,x,y,w,h,fill,ln=None,lw=0.75,shape=MSO_SHAPE.RECTANGLE):
    sp=s.shapes.add_shape(shape,Inches(x),Inches(y),Inches(w),Inches(h))
    sp.fill.solid(); sp.fill.fore_color.rgb=fill; noshadow(sp)
    if ln is None: noline(sp)
    else: sp.line.color.rgb=ln; sp.line.width=Pt(lw)
    return sp

def set_dash(conn,dash='dash'):
    ln=conn.line._get_or_add_ln()
    d=ln.find(qn('a:prstDash'))
    if d is None:
        d=ln.makeelement(qn('a:prstDash'),{}); ln.append(d)
    d.set('val',dash)

def line(s,x1,y1,x2,y2,color,w=1.0,dash=None):
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    c.line.color.rgb=color; c.line.width=Pt(w); noshadow(c)
    if dash: set_dash(c,dash)
    return c

def brand(s,ctx,pageno):
    text(s,0.7,0.34,4,0.4,'TUUCI',18,NAVY,False,'l',SERIF)
    # arc under wordmark
    arc=s.shapes.add_shape(MSO_SHAPE.ARC,Inches(0.72),Inches(0.74),Inches(0.95),Inches(0.18))
    text(s,7.0,0.36,5.6,0.35,ctx,10.5,MUTED,False,'r',BODY)
    text(s,12.4,7.04,0.7,0.3,str(pageno),10,RGBColor.from_string('AAB0B3'),False,'r')
    rect(s,0,0,0.09,7.5,GOLD)

def eyebrow(s,x,y,t): text(s,x,y,11,0.3,t.upper(),11.5,GOLDD,True,'l')

def fnum(v):
    v=round(v,1)
    return ('%g'%v)

# ---------- chart: pareto ----------
def pareto(s,x0,y0,w,h):
    data=sorted(STEPS,key=lambda d:-d[2]); n=len(data); maxT=data[0][2]; total=sum(d[2] for d in data)
    base=y0+h-0.55; top=y0+0.25; ph=base-top; bw=w/n; barW=bw*0.66
    line(s,x0,base,x0+w,base,LINE,1.0)
    cum=0; pts=[]
    for i,(num,name,t) in enumerate(data):
        x=x0+i*bw+(bw-barW)/2; bh=t/maxT*ph; by=base-bh
        over=t>TAKT
        rect(s,x,by,barW,bh,RUST if over else NAVY)
        text(s,x-0.1,by-0.26,barW+0.2,0.24,fnum(t),9,RGBColor.from_string('3A4148'),True,'c')
        text(s,x-0.12,base+0.06,barW+0.24,0.24,'#%d'%num,8.5,MUTED,False,'c')
        cum+=t; pts.append((x0+i*bw+bw/2, base-cum/total*ph))
    ty=base-TAKT/maxT*ph
    line(s,x0,ty,x0+w,ty,RUST,1.4,'dash')
    text(s,x0+w-1.5,ty-0.26,1.5,0.22,'takt 42 min',9,RUST,True,'r')
    for i in range(len(pts)-1):
        line(s,pts[i][0],pts[i][1],pts[i+1][0],pts[i+1][1],GOLD,1.5)

# ---------- chart: yamazumi ----------
def yama(s,x0,y0,w,h,stations,maxY):
    base=y0+h-0.42; top=y0+0.28; ph=base-top; n=len(stations); bw=w/n; barW=min(0.95,bw*0.62)
    line(s,x0,base,x0+w,base,LINE,1.0)
    for idx,(lbl,items) in enumerate(stations):
        cx=x0+idx*bw+bw/2; x=cx-barW/2; yc=base; load=sum(it[1] for it in items)
        over=load>TAKT+1e-6
        for (ilab,it_t,cstep) in items:
            hi=it_t/maxY*ph; y=yc-hi
            rect(s,x,y,barW,hi,RGBColor.from_string(COL.get(cstep,'888888')))
            if hi>=0.22:
                text(s,x,y+hi/2-0.12,barW,0.24,str(ilab),9,WHITE,True,'c')
            yc=y
        text(s,cx-0.6,yc-0.26,1.2,0.22,fnum(load),9,(RUST if over else NAVY),True,'c')
        text(s,cx-0.7,base+0.05,1.4,0.22,lbl,8.5,RGBColor.from_string('3A4148'),False,'c')
    ty=base-TAKT/maxY*ph
    line(s,x0,ty,x0+w,ty,RUST,1.4,'dash')
    text(s,x0+w-1.3,ty-0.24,1.3,0.2,'takt 42',9,RUST,True,'r')

# =================================================================
# SLIDE 1 — TITLE
# =================================================================
s=slide(); rect(s,0,0,0.12,7.5,GOLD)
text(s,1.0,2.0,8,0.7,'TUUCI',40,NAVY,False,'l',SERIF)
ar=s.shapes.add_shape(MSO_SHAPE.ARC,Inches(1.05),Inches(2.78),Inches(2.0),Inches(0.32));
text(s,1.0,3.25,10,0.35,'STANDARD WORK · LINE-BALANCE STUDY',12,GOLDD,True,'l')
text(s,1.0,3.7,11.5,1.6,[('Meritage 3-Seater',44,NAVY,True,BODY)],align='l')
para_tb=text(s,1.0,4.55,11.5,0.8,'Where the labor goes — and how we level it.',24,INK,False,'l')
rect(s,1.02,5.5,1.0,0.05,GOLD)
text(s,1.0,5.7,10.5,0.6,'Findings from the standard-work data, with a proposed balanced line to meet a 10-unit/day demand.',14.5,MUTED,False,'l')
text(s,12.4,7.04,0.7,0.3,'1',10,RGBColor.from_string('AAB0B3'),False,'r')

# =================================================================
# SLIDE 2 — HEADLINES
# =================================================================
s=slide(); brand(s,'Meritage 3-Seater · Headlines',2)
eyebrow(s,0.7,1.15,'The numbers at a glance')
text(s,0.7,1.5,12,0.6,'One Meritage carries 254 minutes of hands-on work',27,NAVY,True,'l')
text(s,0.7,2.2,12,0.4,'Summed straight from the standard work instruction — 11 timed steps (the PPE step has no time).',14,MUTED,False,'l')
cards=[('254','min','Total labor per unit (4.24 labor-hours of hands-on work)',GOLD),
       ('11','','Timed work steps (SWI steps 2–12)',GOLD),
       ('64','min','Biggest single step — Step 4, Arms Assembly (25% of the build)',RUST),
       ('61','%','of all labor sits in just 3 steps (Arms, Seat Frame, Back Frame)',RUST)]
cw=2.78; gap=0.32; x=0.7; y=2.95
for v,u,l,acc in cards:
    rect(s,x,y,cw,1.7,WHITE,LINE,0.75,MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s,x,y,cw,0.07,acc)
    tb=s.shapes.add_textbox(Inches(x+0.18),Inches(y+0.22),Inches(cw-0.36),Inches(0.6)); tf=tb.text_frame; tf.word_wrap=True
    p=tf.paragraphs[0]; r=p.add_run(); r.text=v; r.font.size=Pt(32); r.font.bold=True; r.font.name=BODY; r.font.color.rgb=NAVY
    if u: ru=p.add_run(); ru.text=' '+u; ru.font.size=Pt(13); ru.font.bold=True; ru.font.color.rgb=MUTED; ru.font.name=BODY
    text(s,x+0.18,y+0.92,cw-0.36,0.7,l,11,RGBColor.from_string('4A4F53'),False,'l')
    x+=cw+gap
rect(s,0.7,5.0,11.93,1.15,RGBColor.from_string('FFF7EC'),RGBColor.from_string('F0DCBD'),0.75,MSO_SHAPE.ROUNDED_RECTANGLE)
rect(s,0.7,5.0,0.07,1.15,RUST)
tb=text(s,1.0,5.22,11.3,0.9,[('Labor-hours vs. clock time.  ',13.5,NAVY,True,BODY),
    ('254 min is total hands-on work — if one person built it alone. With a few builders working in parallel the wall-clock time per unit is much shorter; this study is about ',13.5,INK,False,BODY),
    ('spreading those 254 minutes evenly',13.5,NAVY,True,BODY),(' so the line flows.',13.5,INK,False,BODY)])

# =================================================================
# SLIDE 3 — BOTTLENECKS
# =================================================================
s=slide(); brand(s,'Meritage 3-Seater · Bottlenecks',3)
eyebrow(s,0.7,1.15,'Where the time concentrates')
text(s,0.7,1.5,12,0.6,'Three assemblies dominate — and each is bigger than takt',27,NAVY,True,'l')
pareto(s,0.7,2.35,7.2,3.9)
text(s,0.7,6.35,7.2,0.4,[('Bars = minutes per step (sorted).  ',11,MUTED,False,BODY),('Gold line = cumulative %.  ',11,GOLDD,True,BODY),('Dashed = 42-min takt.',11,RUST,True,BODY)])
bx=8.2
tb=text(s,bx,2.5,4.4,0.5,[('Step 4 — Arms Assembly: 64 min.',14.5,NAVY,True,BODY)])
para(tb,[(' A quarter of the build, and 1.5× the takt by itself.',14.5,INK,False,BODY)],space_before=2)
tb2=text(s,bx,3.5,4.4,0.5,[('Step 5 — Seat Frame: 47 min',14.5,NAVY,True,BODY),(' and ',14.5,INK,False,BODY),('Step 6 — Back Frame: 43.5 min.',14.5,NAVY,True,BODY)])
tb3=text(s,bx,4.55,4.4,0.5,[('These three steps = 61% of the labor.',14.5,NAVY,True,BODY),(' Fix these and you move the line.',14.5,INK,False,BODY)])
tb4=text(s,bx,5.5,4.4,0.7,[('All three exceed the 42-min takt,',14.5,NAVY,True,BODY),(' so alone they cap output no matter how the rest is staffed.',14.5,INK,False,BODY)])
text(s,0.7,6.75,11.9,0.3,'Takt = 420 min (7 hr) ÷ 10 units = 42 min per unit — the pace we must match.',12,MUTED,False,'l')

# =================================================================
# SLIDE 4 — OPPORTUNITY
# =================================================================
s=slide(); brand(s,'Meritage 3-Seater · The Opportunity',4)
eyebrow(s,0.7,1.15,'Where we can help')
text(s,0.7,1.5,12,0.6,'Put hands where the hours are',27,NAVY,True,'l')
text(s,0.7,2.2,11.8,0.5,'The three big assemblies are buildable in parallel — more than one person can work them at once. Splitting them is what brings every station under takt.',14,MUTED,False,'l')
# table
tbl=s.shapes.add_table(4,3,Inches(0.7),Inches(3.0),Inches(5.6),Inches(2.2)).table
tbl.columns[0].width=Inches(3.2); tbl.columns[1].width=Inches(1.2); tbl.columns[2].width=Inches(1.2)
hdr=['Step','Now','2 builders']
rows=[('#4  Arms Assembly','64','32'),('#5  Seat Frame Assembly','47','23.5'),('#6  Back Frame Assembly','43.5','21.8')]
for c,htext in enumerate(hdr):
    cell=tbl.cell(0,c); cell.fill.solid(); cell.fill.fore_color.rgb=NAVY
    cell.text=htext; pr=cell.text_frame.paragraphs[0]; pr.runs[0].font.size=Pt(11); pr.runs[0].font.bold=True; pr.runs[0].font.color.rgb=WHITE; pr.runs[0].font.name=BODY
    if c>0: pr.alignment=PP_ALIGN.RIGHT
for r,(a,b,c) in enumerate(rows,1):
    for ci,val in enumerate((a,b,c)):
        cell=tbl.cell(r,ci); cell.fill.solid(); cell.fill.fore_color.rgb=WHITE
        cell.text=val; pp=cell.text_frame.paragraphs[0]; rn=pp.runs[0]
        rn.font.size=Pt(12); rn.font.name=BODY; rn.font.color.rgb=(GOOD if ci==2 else NAVY if ci==0 else INK); rn.font.bold=(ci>0)
        if ci>0: pp.alignment=PP_ALIGN.RIGHT
text(s,0.7,5.35,5.6,0.5,'Minutes per unit. Splitting helps only where the work is genuinely separable — these assemblies are.',11.5,MUTED,False,'l')
bx=6.9
text(s,bx,3.0,5.7,0.6,[('The principle: ',14,NAVY,True,BODY),('a step’s labor doesn’t disappear when you split it — but the time on the clock drops, so it stops gating the line.',14,INK,False,BODY)])
text(s,bx,4.1,5.7,0.5,[('Arms (64 → 32)',14,NAVY,True,BODY),(' alone takes the worst bottleneck below takt.',14,INK,False,BODY)])
text(s,bx,4.85,5.7,0.6,[('Smaller steps (legs, attach, finishing, plate) are grouped to fill the remaining operators evenly.',14,INK,False,BODY)])
text(s,bx,5.75,5.7,0.5,[('Net effect (next slide): ',14,NAVY,True,BODY),('a level line that hits 10/day.',14,GOOD,True,BODY)])

# =================================================================
# SLIDE 5 — THE BALANCE
# =================================================================
s=slide(); brand(s,'Meritage 3-Seater · The Balance',5)
eyebrow(s,0.7,1.15,'The line balance')
text(s,0.7,1.5,12,0.6,'From lumpy to level',27,NAVY,True,'l')
text(s,0.7,2.25,5.8,0.4,'Today — work as it sits',14,RUST,True,'l')
text(s,6.95,2.25,5.8,0.4,'Proposed — balanced to takt',14,GOOD,True,'l')
NOW=[('#%d'%n,[(n,t,n)]) for (n,name,t) in STEPS]
PROP=[('S1',[(2,19.25,2),(3,8,3)]),
      ('S2 ×2',[(4,32,4)]),
      ('S3',[('5a',41,5)]),
      ('S4',[('5b',6,5),('6a',33.5,6)]),
      ('S5',[('6b',10,6),(7,18,7),(8,5,8),(9,8,9)]),
      ('S6',[(10,22.5,10),(11,18,11),(12,1,12)])]
yama(s,0.7,2.65,5.7,3.4,NOW,68)
yama(s,6.95,2.65,5.7,3.4,PROP,68)
text(s,0.7,6.15,5.7,0.6,'One operator per step: Arms towers at 64 min → a unit only every 64 min (~6.6/day), and small steps sit idle.',11,MUTED,False,'l')
text(s,6.95,6.15,5.7,0.6,'7 operators across 6 stations (Arms built 2-up; Seat & Back Frame split). Every station under takt → ~10.1/day, 88% utilized.',11,MUTED,False,'l')

# =================================================================
# SLIDE 6 — HOW IT LOOKS
# =================================================================
s=slide(); brand(s,'Meritage 3-Seater · How it looks',6)
eyebrow(s,0.7,1.15,'How it would actually look on the floor')
text(s,0.7,1.5,12,0.6,'7 operators across 6 stations — every station under takt',27,NAVY,True,'l')
tbl=s.shapes.add_table(7,3,Inches(0.7),Inches(2.4),Inches(8.7),Inches(3.9)).table
tbl.columns[0].width=Inches(1.7); tbl.columns[1].width=Inches(6.0); tbl.columns[2].width=Inches(1.0)
for c,htext in enumerate(['Station','Work (SWI steps)','Min/unit']):
    cell=tbl.cell(0,c); cell.fill.solid(); cell.fill.fore_color.rgb=NAVY; cell.text=htext
    pr=cell.text_frame.paragraphs[0]; pr.runs[0].font.size=Pt(11); pr.runs[0].font.bold=True; pr.runs[0].font.color.rgb=WHITE; pr.runs[0].font.name=BODY
    if c==2: pr.alignment=PP_ALIGN.RIGHT
strows=[('S1','#2 Connector Prep · #3 Leg Assembly','27.3'),
        ('S2 · 2 builders','#4 Arms Assembly (built together)','32.0'),
        ('S3','#5 Seat Frame Assembly (first part)','41.0'),
        ('S4','#5 Seat Frame (rest) · #6 Back Frame (first part)','39.5'),
        ('S5','#6 Back Frame (rest) · #7 Attaching Arms · #8 Back Support · #9 Leg Finishing','41.0'),
        ('S6','#10 Seat Support · #11 Attaching Seat Frame · #12 TUUCI Plate','41.5')]
for r,(a,b,c) in enumerate(strows,1):
    for ci,val in enumerate((a,b,c)):
        cell=tbl.cell(r,ci); cell.fill.solid(); cell.fill.fore_color.rgb=WHITE; cell.text=val
        pp=cell.text_frame.paragraphs[0]; rn=pp.runs[0]; rn.font.size=Pt(11); rn.font.name=BODY
        rn.font.color.rgb=(NAVY if ci!=1 else INK); rn.font.bold=(ci!=1)
        if ci==2: pp.alignment=PP_ALIGN.RIGHT
# stat cards right
sx=9.7
for i,(v,u,l,acc) in enumerate([('10.1','/day','capacity vs. 10/day demand — met',GOOD),
                                ('41.5','min','slowest station (S6) — under takt',GOLD),
                                ('88','%','line efficiency (very little idle)',GOOD)]):
    yy=2.4+i*1.32
    rect(s,sx,yy,2.95,1.18,WHITE,LINE,0.75,MSO_SHAPE.ROUNDED_RECTANGLE); rect(s,sx,yy,2.95,0.06,acc)
    tb=s.shapes.add_textbox(Inches(sx+0.16),Inches(yy+0.16),Inches(2.6),Inches(0.5)); tf=tb.text_frame
    p=tf.paragraphs[0]; rr=p.add_run(); rr.text=v; rr.font.size=Pt(26); rr.font.bold=True; rr.font.name=BODY; rr.font.color.rgb=NAVY
    if u: ru=p.add_run(); ru.text=' '+u; ru.font.size=Pt(12); ru.font.bold=True; ru.font.color.rgb=MUTED; ru.font.name=BODY
    text(s,sx+0.16,yy+0.7,2.6,0.4,l,10.5,RGBColor.from_string('4A4F53'),False,'l')
text(s,0.7,6.5,11.9,0.4,'S2 is two builders working Arms together (64 → 32 min). Seat Frame (#5) and Back Frame (#6) are each split across two stations to level the load. Times measured from the SWI.',11.5,MUTED,False,'l')

# =================================================================
# SLIDE 7 — SUMMARY
# =================================================================
s=slide(); brand(s,'Meritage 3-Seater · Summary',7)
eyebrow(s,0.7,1.15,'What this means')
text(s,0.7,1.5,12,0.6,'Same product, leveled to demand',27,NAVY,True,'l')
bullets=[('254 min of labor',', with 61% in three assemblies — Arms (#4), Seat Frame (#5), Back Frame (#6).'),
         ('Those three each run longer than takt',', so they’re the constraint.'),
         ('Build Arms 2-up, split Seat & Back Frame, group the rest',' → 7 operators across 6 stations, every station ≤ 41.5 min.'),
         ('Result: ~10.1 units/day at 88% efficiency',', meeting the 10/day demand.')]
y=2.5
for b,rest in bullets:
    rect(s,0.75,y+0.07,0.12,0.12,GOLD,shape=MSO_SHAPE.OVAL)
    text(s,1.05,y-0.05,5.6,0.8,[(b,14.5,NAVY,True,BODY),(rest,14.5,INK,False,BODY)],line_spacing=1.05)
    y+=0.95
# caveat boxes
cx=7.0
rect(s,cx,2.45,5.6,1.7,RGBColor.from_string('FFF7EC'),RGBColor.from_string('F0DCBD'),0.75,MSO_SHAPE.ROUNDED_RECTANGLE); rect(s,cx,2.45,0.07,1.7,RUST)
text(s,cx+0.25,2.62,5.2,1.5,[('Assumptions & caveats.  ',13,NAVY,True,BODY),('Demand 10/day over a 7-hour day → 42-min takt. Times are measured from the SWI. Splitting helps only where a task is genuinely separable; sequence-locked work won’t overlap — to confirm on the floor.',13,RGBColor.from_string('3A2F20'),False,BODY)],line_spacing=1.08)
rect(s,cx,4.35,5.6,1.5,RGBColor.from_string('EEF3EE'),RGBColor.from_string('CFE0D3'),0.75,MSO_SHAPE.ROUNDED_RECTANGLE); rect(s,cx,4.35,0.07,1.5,GOOD)
text(s,cx+0.25,4.52,5.2,1.3,[('Next step.  ',13,NAVY,True,BODY),('Validate the three split assemblies on the floor, then trial the 6-station layout. The analyzer re-balances live as real times firm up.',13,RGBColor.from_string('23402C'),False,BODY)],line_spacing=1.08)

prs.save('standalone/Meritage_Line_Balance.pptx')
print('saved standalone/Meritage_Line_Balance.pptx  (%d slides)'%len(prs.slides._sldIdLst))
